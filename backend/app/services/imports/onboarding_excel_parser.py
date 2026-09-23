from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal
from io import BytesIO
from typing import Any
from zipfile import BadZipFile

from openpyxl import load_workbook
from openpyxl.utils.exceptions import (
    InvalidFileException,
)

from app.exceptions.base import BusinessRuleError
from app.services.imports.onboarding_excel_columns import (
    ONBOARDING_COLUMN_COUNT,
    ONBOARDING_COLUMN_KEYS,
    ONBOARDING_SHEET_NAME,
)


DATE_KEYS = {
    "hire_date",
    "introductory_internship_date",
    "training_date",
    "test_date",
    "main_internship_start_date",
    "main_internship_end_date",
    "first_payment_due_date",
    "first_payment_created_at",
    "second_payment_due_date",
    "second_payment_created_at",
    "stage_1_planned_start_date",
    "stage_1_planned_end_date",
    "stage_1_actual_date",
    "stage_2_planned_start_date",
    "stage_2_planned_end_date",
    "stage_2_actual_date",
    "stage_3_planned_end_date",
    "stage_3_actual_date",
    "separation_date",
}


BOOLEAN_KEYS = {
    "internship_form_completed",
    "mentor_payment_expected",
    "payment_is_fully_paid",
    "stage_1_deadline_compliant",
    "stage_2_deadline_compliant",
    "stage_3_deadline_compliant",
    "probation_completed_successfully",
}


INTEGER_KEYS = {
    "source_id",
    "actual_internship_duration_days",
    "days_since_hire",
    "days_since_stage_1",
}


MONEY_KEYS = {
    "first_recommended_amount",
    "first_actual_amount",
    "second_recommended_amount",
    "second_actual_amount",
    "total_recommended_amount",
    "total_actual_amount",
}


@dataclass(frozen=True)
class ParsedOnboardingRow:
    excel_row_number: int
    source_row_id: str | None
    raw_data: dict[str, Any]
    normalized_data: dict[str, Any]
    warnings: list[str]
    errors: list[str]


def parse_onboarding_excel(
    content: bytes,
) -> list[ParsedOnboardingRow]:
    if not content:
        raise BusinessRuleError(
            "The uploaded Excel file is empty"
        )

    try:
        workbook = load_workbook(
            filename=BytesIO(content),
            read_only=True,
            data_only=True,
            keep_links=False,
        )
    except (
        BadZipFile,
        InvalidFileException,
        OSError,
        ValueError,
    ) as error:
        raise BusinessRuleError(
            "The uploaded file is not a valid "
            "Excel workbook"
        ) from error

    try:
        if ONBOARDING_SHEET_NAME not in workbook.sheetnames:
            raise BusinessRuleError(
                f'Worksheet "{ONBOARDING_SHEET_NAME}" '
                "was not found"
            )

        worksheet = workbook[
            ONBOARDING_SHEET_NAME
        ]

        header_row_number = _find_header_row(
            worksheet
        )

        adaptation_columns = (
            _find_adaptation_columns(
                worksheet=worksheet,
                header_row_number=header_row_number,
            )
        )

        parsed_rows: list[
            ParsedOnboardingRow
        ] = []

        for row_number, values in enumerate(
            worksheet.iter_rows(
                min_row=header_row_number + 1,
                min_col=1,
                max_col=ONBOARDING_COLUMN_COUNT,
                values_only=True,
            ),
            start=header_row_number + 1,
        ):
            if _is_empty_row(values):
                continue

            parsed_rows.append(
                _parse_row(
                    row_number=row_number,
                    values=values,
                    overrides={
                        key: values[
                            column_number - 1
                            ]
                        for key, column_number
                        in adaptation_columns.items()
                    },
                )
            )

        if not parsed_rows:
            raise BusinessRuleError(
                'Worksheet "ВА" does not contain '
                "any data rows"
            )

        return parsed_rows
    finally:
        workbook.close()


def _find_header_row(worksheet) -> int:
    for row_number in range(
        1,
        min(worksheet.max_row, 10) + 1,
    ):
        first_value = _normalize_text(
            worksheet.cell(
                row=row_number,
                column=1,
            ).value
        )

        second_value = _normalize_text(
            worksheet.cell(
                row=row_number,
                column=2,
            ).value
        )

        if (
            first_value.upper() == "ID"
            and second_value.upper() == "РЦ"
        ):
            return row_number

    raise BusinessRuleError(
        'Header row was not found in worksheet "ВА"'
    )


def _parse_row(
    row_number: int,
    values: tuple[Any, ...],
    overrides: dict[str, Any] | None = None,
) -> ParsedOnboardingRow:
    source_data = dict(
        zip(
            ONBOARDING_COLUMN_KEYS,
            values,
            strict=True,
        )
    )

    if overrides:
        source_data.update(overrides)

    raw_data = {
        key: _json_value(value)
        for key, value
        in source_data.items()
    }

    warnings: list[str] = []
    errors: list[str] = []

    normalized_data = {
        key: _normalize_value(
            key=key,
            value=value,
            warnings=warnings,
        )
        for key, value
        in source_data.items()
    }

    employee_name = normalized_data[
        "employee_name"
    ]

    if not employee_name:
        errors.append(
            "Не указано ФИО сотрудника"
        )

    if not normalized_data[
        "distribution_center"
    ]:
        errors.append(
            "Не указан распределительный центр"
        )

    if not normalized_data[
        "distribution_center_division"
    ]:
        errors.append(
            "Не указано подразделение РЦ"
        )

    if not normalized_data["position_name"]:
        errors.append(
            "Не указана должность"
        )

    if not normalized_data["tutor_name"]:
        errors.append(
            "Не указан куратор МПО"
        )

    if not normalized_data["personnel_number"]:
        warnings.append(
            "Не указан табельный номер"
        )

    training_date = normalized_data.get(
        "training_date"
    )

    test_date = normalized_data.get(
        "test_date"
    )

    if (
            training_date is not None
            and test_date is not None
            and test_date < training_date
    ):
        errors.append(
            "Дата проверки не может быть "
            "раньше даты обучения"
        )

    source_id_value = normalized_data[
        "source_id"
    ]

    return ParsedOnboardingRow(
        excel_row_number=row_number,
        source_row_id=(
            str(source_id_value)
            if source_id_value is not None
            else None
        ),
        raw_data=raw_data,
        normalized_data=normalized_data,
        warnings=warnings,
        errors=errors,
    )


def _normalize_value(
    key: str,
    value: Any,
    warnings: list[str],
) -> Any:
    if value is None:
        return None

    if key in DATE_KEYS:
        return _normalize_date(
            key=key,
            value=value,
            warnings=warnings,
        )

    if key in BOOLEAN_KEYS:
        return _normalize_boolean(
            key=key,
            value=value,
            warnings=warnings,
        )

    if key in INTEGER_KEYS:
        return _normalize_integer(
            key=key,
            value=value,
            warnings=warnings,
        )

    if key in MONEY_KEYS:
        return _normalize_money(
            key=key,
            value=value,
            warnings=warnings,
        )

    if key == "test_result":
        return _normalize_test_result(
            value=value,
            warnings=warnings,
        )

    if isinstance(value, str):
        return _normalize_text(value) or None

    return value


def _normalize_date(
    key: str,
    value: Any,
    warnings: list[str],
) -> str | None:
    if isinstance(value, datetime):
        return value.date().isoformat()

    if isinstance(value, date):
        return value.isoformat()

    if isinstance(value, str):
        normalized = _normalize_text(value)

        if not normalized:
            return None

        for date_format in (
            "%d.%m.%Y",
            "%Y-%m-%d",
            "%d/%m/%Y",
        ):
            try:
                return datetime.strptime(
                    normalized,
                    date_format,
                ).date().isoformat()
            except ValueError:
                continue

    warnings.append(
        f'Field "{key}" contains an invalid date'
    )

    return None


def _normalize_boolean(
    key: str,
    value: Any,
    warnings: list[str],
) -> bool | None:
    if isinstance(value, bool):
        return value

    if isinstance(value, (int, float)):
        if value in (0, 1):
            return bool(value)

    normalized = _normalize_text(value).lower()

    if normalized in {
        "да",
        "yes",
        "true",
        "1",
    }:
        return True

    if normalized in {
        "нет",
        "no",
        "false",
        "0",
    }:
        return False

    warnings.append(
        f'Field "{key}" contains an invalid '
        "boolean value"
    )

    return None


def _normalize_integer(
    key: str,
    value: Any,
    warnings: list[str],
) -> int | None:
    try:
        if isinstance(value, str):
            value = value.replace(
                " ",
                "",
            ).replace(
                ",",
                ".",
            )

        return int(float(value))
    except (TypeError, ValueError):
        warnings.append(
            f'Field "{key}" contains an invalid number'
        )

        return None


def _normalize_money(
    key: str,
    value: Any,
    warnings: list[str],
) -> str | None:
    try:
        if isinstance(value, str):
            value = (
                value
                .replace("₽", "")
                .replace(" ", "")
                .replace(",", ".")
            )

        return str(
            Decimal(str(value)).quantize(
                Decimal("0.01")
            )
        )
    except Exception:
        warnings.append(
            f'Field "{key}" contains an invalid amount'
        )

        return None


def _normalize_test_result(
    value: Any,
    warnings: list[str],
) -> float | None:
    try:
        if isinstance(value, str):
            normalized = (
                value
                .replace("%", "")
                .replace(" ", "")
                .replace(",", ".")
            )

            result = float(normalized)
        else:
            result = float(value)

        if 0 <= result <= 1:
            result *= 100

        if result < 0 or result > 100:
            raise ValueError

        return round(result, 2)
    except (TypeError, ValueError):
        warnings.append(
            'Field "test_result" contains an invalid '
            "percentage"
        )

        return None


def _normalize_text(value: Any) -> str:
    if value is None:
        return ""

    return " ".join(
        str(value).replace(
            "\xa0",
            " ",
        ).split()
    )


def _json_value(value: Any) -> Any:
    if isinstance(value, datetime):
        return value.isoformat()

    if isinstance(value, date):
        return value.isoformat()

    if isinstance(value, Decimal):
        return str(value)

    return value


def _is_empty_row(
    values: tuple[Any, ...],
) -> bool:
    return all(
        value is None
        or (
            isinstance(value, str)
            and not value.strip()
        )
        for value in values
    )


def _find_adaptation_columns(
    worksheet,
    header_row_number: int,
) -> dict[str, int]:
    expected_headers = {
        "зона 1": "stage_1_zone",
        "риск зоны 1": "stage_1_risk_zone",
        "причина риска 1": "stage_1_risk_reason",

        "зона 2": "stage_2_zone",
        "риск зоны 2": "stage_2_risk_zone",
        "причина риска 2": "stage_2_risk_reason",

        "зона 3": "stage_3_zone",
        "риск зоны 3": "stage_3_risk_zone",
        "причина риска 3": "stage_3_risk_reason",
    }

    headers = next(
        worksheet.iter_rows(
            min_row=header_row_number,
            max_row=header_row_number,
            values_only=True,
        )
    )

    result: dict[str, int] = {}

    for column_number, value in enumerate(
        headers,
        start=1,
    ):
        header_key = (
            _normalize_text(value)
            .replace("_", "")
            .casefold()
        )

        field_name = expected_headers.get(
            header_key
        )

        if field_name is not None:
            result[field_name] = (
                column_number
            )

    return result