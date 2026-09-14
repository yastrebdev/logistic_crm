from collections.abc import Sequence
from datetime import date
from io import BytesIO
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import (
    Alignment,
    Font,
    PatternFill,
)
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import (
    Table,
    TableStyleInfo,
)

from app.schemas.onboarding_analytics import (
    OnboardingAnalyticsRow,
)


HEADERS = [
    "ID процесса",
    "РЦ",
    "Подразделение",
    "Группа подразделения",
    "Руководитель",
    "Табельный номер",
    "ФИО",
    "Должность",
    "Статус должности",
    "Дата ТУ",
    "Причина не ТУ",
    "ТУ не сразу (причина)",
    "Комментарий МПО",
    "Статус ТУ",
    "Начало вводного процесса",
    "Завершение вводного процесса",
    "Дата ознакомительной стажировки",
    "Наставник ознакомительной стажировки",
    "Куратор МПО",
    "Дата обучения",
    "Формат допуска на должность",
    "Дата проверки",
    "Результат проверки",
    "Дата начала основной стажировки",
    "Дата окончания основной стажировки",
    "Продолжительность стажировки",
    "Количество дней стажировки",
    "Срок стажировки соблюдён",
    "Наставник основной стажировки",
    "Должность наставника",
    "Статус назначения наставника",
    "Лист стажировки получен",
    "Наставнику предусмотрена оплата",
    "Срок заведения оплаты",
    "Рекомендуемая сумма выплаты",
    "Фактическая сумма выплаты",
    "Дата заведения оплаты",
    "Дата выплаты",
    "Статус оплаты",
    "Оплата произведена полностью",
    "Причина неполной оплаты",
    "Как заведена премия в 1С",
    "Количество дней после ТУ",
    "Статус процесса адаптации",
    "Общий срок адаптации",
    "1 адаптация с",
    "1 адаптация до",
    "Дата 1 адаптации",
    "Метод проведения 1 адаптации",
    "Срок 1 адаптации соблюдён",
    "Участники 1 адаптации",
    "Причина опоздания 1 адаптации",
    "Риск зоны 1",
    "Причина риска 1",
    "Комментарий МПО 1",
    "Количество дней после 1 адаптации",
    "2 адаптация с",
    "2 адаптация до",
    "Дата 2 адаптации",
    "Метод проведения 2 адаптации",
    "Срок 2 адаптации соблюдён",
    "Участники 2 адаптации",
    "Причина опоздания 2 адаптации",
    "Риск зоны 2",
    "Причина риска 2",
    "Комментарий МПО 2",
    "3 адаптация с",
    "3 адаптация до",
    "Дата 3 адаптации",
    "Метод проведения 3 адаптации",
    "Срок 3 адаптации соблюдён",
    "Участники 3 адаптации",
    "Причина опоздания 3 адаптации",
    "Риск зоны 3",
    "Причина риска 3",
    "Комментарий МПО 3",
    "ИС пройден успешно",
    "Дата увольнения",
    "Причина увольнения",
    "Причина увольнения — ОС руководителя",
]


LABELS = {
    # Статус должности
    "line_staff": "Линейный персонал",
    "specialist": "Специалист",
    "manager": "Руководитель",
    "head": "Директор",

    # Статус ТУ
    "former_employee": "Работал ранее",
    "external_candidate": "Внешний кандидат",

    # Причина не ТУ
    "candidate_refused": "Отказ кандидата",
    "employer_refused": "Отказ работодателя",
    "document_issues": "Проблемы с документами",
    "medical_restrictions": "Медицинские ограничения",
    "failed_background_check": "Не пройдена проверка",
    "position_closed": "Вакансия закрыта",
    "no_contact": "Нет связи с кандидатом",
    "other": "Другое",

    # Причина задержки ТУ
    "documents_pending": "Ожидание документов",
    "medical_exam_pending": "Ожидание медосмотра",
    "background_check_pending": "Ожидание проверки",
    "candidate_request": "По просьбе кандидата",
    "employer_request": "По инициативе работодателя",
    "start_date_postponed": "Перенос даты выхода",

    # Формат допуска
    "The test is on the form": "Тест на бланке",
    "The test on the portal": "Тест на портале",
    "Exam": "Экзамен",
    "Interview": "Собеседование",

    # Назначение наставника
    "Previously employed": "Работал ранее",
    "MPO is the mentor": "МПО — наставник",
    "Mentor is not required": "Наставник не предполагается",
    "Mentor is assigned": "Наставник назначен",

    # Оплата
    "Pending": "Ожидает обработки",
    "Approved": "Одобрена",
    "Paid": "Выплачена",
    "Cancelled": "Отменена",

    "Internship not completed":
        "Стажировка не завершена",
    "Insufficient internship duration":
        "Недостаточная продолжительность стажировки",
    "Mentor is not eligible for payment":
        "Наставнику не предусмотрена выплата",
    "Employee left before completion":
        "Сотрудник уволился до завершения",
    "Mentor left before completion":
        "Наставник уволился до завершения",
    "Duplicate payment":
        "Дублирующая выплата",
    "Payment already processed":
        "Выплата уже обработана",
    "Incorrect data":
        "Некорректные данные",
    "Management decision":
        "Решение руководства",
    "Other": "Другое",

    # Процесс адаптации
    "Active": "В процессе",
    "Completed": "Завершена",
    "Cancelled": "Отменена",

    # Метод адаптации
    "in person": "Очно",
    "on portal": "На портале",
    "a phone call": "Телефонный звонок",
    "yandex form": "Яндекс-форма",

    # Участники
    "MPO": "МПО",
    "MPO and MPP": "МПО, МПП",
    "MPO and supervisor": "МПО, руководитель",
    "MPO, MPP and supervisor":
        "МПО, МПП, руководитель",

    # Причины опоздания
    "Employee absent": "Отсутствие сотрудника",
    "Supervisor absent": "Отсутствие руководителя",
    "MPO absent": "Отсутствие МПО",
    "MPP absent": "Отсутствие МПП",
    "Schedule conflict": "Конфликт расписания",
    "Technical issues": "Технические проблемы",
    "High workload": "Высокая рабочая нагрузка",
    "Adaptation rescheduled": "Адаптация перенесена",

    # Риск
    "Green": "Зелёный",
    "Yellow": "Жёлтый",
    "Red": "Красный",

    "Low performance": "Низкая производительность",
    "Insufficient skills": "Недостаточные навыки",
    "Low motivation": "Низкая мотивация",
    "Attendance issues": "Проблемы с посещаемостью",
    "Disciplinary issues": "Дисциплинарные нарушения",
    "Difficulties working with the team":
        "Сложности в работе с коллективом",
    "Difficulties working with the supervisor":
        "Сложности в работе с руководителем",
    "Failure to meet adaptation goals":
        "Невыполнение целей адаптации",

    # Увольнение
    "voluntary_resignation":
        "По собственному желанию",
    "employer_termination":
        "По инициативе работодателя",
    "mutual_agreement":
        "По соглашению сторон",
    "end_of_contract":
        "Окончание договора",
    "transfer": "Перевод",
    "retirement": "Выход на пенсию",
    "job_abandonment":
        "Неявка на работу",
}


DATE_HEADERS = {
    "Дата ТУ",
    "Начало вводного процесса",
    "Завершение вводного процесса",
    "Дата ознакомительной стажировки",
    "Дата обучения",
    "Дата проверки",
    "Дата начала основной стажировки",
    "Дата окончания основной стажировки",
    "Срок заведения оплаты",
    "Дата заведения оплаты",
    "Дата выплаты",
    "Общий срок адаптации",
    "1 адаптация с",
    "1 адаптация до",
    "Дата 1 адаптации",
    "2 адаптация с",
    "2 адаптация до",
    "Дата 2 адаптации",
    "3 адаптация с",
    "3 адаптация до",
    "Дата 3 адаптации",
    "Дата увольнения",
}


MONEY_HEADERS = {
    "Рекомендуемая сумма выплаты",
    "Фактическая сумма выплаты",
}


PERCENT_HEADERS = {
    "Результат проверки",
}


def _safe_excel_value(
    value: Any,
) -> Any:
    if (
        isinstance(value, str)
        and value.startswith(
            ("=", "+", "-", "@")
        )
    ):
        return f"'{value}"

    return value


def create_onboarding_excel(
    rows: Sequence[OnboardingAnalyticsRow],
) -> BytesIO:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Вводное и адаптация"

    worksheet.append(HEADERS)

    for row in rows:
        worksheet.append(
            [
                _safe_excel_value(value)
                for value in _build_excel_row(
                    row
                )
            ]
        )

    _format_worksheet(
        worksheet=worksheet,
        row_count=len(rows),
    )

    output = BytesIO()
    workbook.save(output)
    output.seek(0)

    return output


def _build_excel_row(
    row: OnboardingAnalyticsRow,
) -> list[Any]:
    stage_1 = row.adaptation_stage_1
    stage_2 = row.adaptation_stage_2
    stage_3 = row.adaptation_stage_3

    return [
        row.introductory_process_id,
        row.distribution_center_name,
        row.division_name,
        row.division_group_abbreviation,
        row.manager_name,
        row.personnel_number,
        row.employee_name,
        row.position_name,
        _label(row.position_category),
        row.hire_date,
        _label(row.reason_not_hiring),
        _label(row.reason_delayed_hiring),
        row.hiring_comment,
        _label(row.candidate_type),
        row.introductory_start_date,
        row.introductory_end_date,
        row.introductory_internship_date,
        row.introductory_mentor_name,
        row.tutor_name or row.tutor_email,
        row.training_date,
        _label(row.admission_format),
        row.test_date,
        _percentage(row.test_result),
        row.main_internship_start_date,
        row.main_internship_end_date,
        _duration_label(row),
        row.actual_internship_duration_days,
        _yes_no(
            row.internship_duration_compliant
        ),
        row.main_mentor_name,
        row.main_mentor_position_name,
        _label(row.mentor_assignment_status),
        _yes_no(
            row.internship_form_completed
        ),
        _yes_no(
            row.mentor_payment_expected
        ),
        row.payment_due_date,
        row.planned_payment_amount,
        row.actual_payment_amount,
        row.payment_created_at,
        row.payment_paid_at,
        _label(row.payment_status),
        _yes_no(
            row.payment_is_fully_paid
        ),
        _label(row.non_payment_reason),
        row.payment_registration_method,
        row.days_since_hire,
        _label(row.adaptation_status),
        row.adaptation_deadline_date,

        *_stage_values(stage_1),

        row.days_since_stage_1,

        *_stage_values(stage_2),
        *_stage_values(stage_3),

        _yes_no(
            row.probation_completed_successfully
        ),
        row.separation_date,
        _label(row.separation_reason),
        row.manager_separation_feedback,
    ]


def _stage_values(stage) -> list[Any]:
    if stage is None:
        return [
            None,
            None,
            None,
            None,
            None,
            None,
            None,
            None,
            None,
            None,
        ]

    deadline_compliant = (
        stage.actual_date is not None
        and stage.actual_date
        <= stage.planned_end_date
    )

    return [
        stage.planned_start_date,
        stage.planned_end_date,
        stage.actual_date,
        _label(stage.method),
        _yes_no(deadline_compliant),
        _label(stage.participants),
        _label(stage.delay_reason),
        _label(stage.risk_zone),
        _label(stage.risk_reason),
        stage.comment,
    ]


def _duration_label(
    row: OnboardingAnalyticsRow,
) -> str | None:
    minimum = (
        row.internship_duration_min_days
    )
    maximum = (
        row.internship_duration_max_days
    )

    if minimum is None or maximum is None:
        return None

    if minimum == maximum:
        return f"{minimum} дней"

    return f"{minimum}–{maximum} дней"


def _label(value) -> str | None:
    if value is None:
        return None

    raw_value = getattr(
        value,
        "value",
        value,
    )

    return LABELS.get(
        str(raw_value),
        str(raw_value),
    )


def _yes_no(
    value: bool | None,
) -> str | None:
    if value is None:
        return None

    return "ДА" if value else "НЕТ"


def _percentage(
    value: float | None,
) -> float | None:
    if value is None:
        return None

    return value / 100


def _format_worksheet(
    worksheet,
    row_count: int,
) -> None:
    header_fill = PatternFill(
        fill_type="solid",
        fgColor="1F4E78",
    )

    header_font = Font(
        color="FFFFFF",
        bold=True,
    )

    for cell in worksheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True,
        )

    worksheet.row_dimensions[1].height = 60
    worksheet.freeze_panes = "G2"
    worksheet.auto_filter.ref = (
        f"A1:{get_column_letter(len(HEADERS))}"
        f"{max(row_count + 1, 1)}"
    )

    if row_count > 0:
        table_reference = (
            f"A1:{get_column_letter(len(HEADERS))}"
            f"{row_count + 1}"
        )

        table = Table(
            displayName="OnboardingAnalytics",
            ref=table_reference,
        )

        table.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
            showColumnStripes=False,
        )

        worksheet.add_table(table)

    for column_index, header in enumerate(
        HEADERS,
        start=1,
    ):
        column_letter = get_column_letter(
            column_index
        )

        worksheet.column_dimensions[
            column_letter
        ].width = _column_width(header)

        for cell in worksheet[
            column_letter
        ][1:]:
            cell.alignment = Alignment(
                vertical="top",
                wrap_text=True,
            )

            if header in DATE_HEADERS:
                cell.number_format = (
                    "dd.mm.yyyy"
                )

            if header in MONEY_HEADERS:
                cell.number_format = (
                    '#,##0.00 "₽"'
                )

            if header in PERCENT_HEADERS:
                cell.number_format = "0.00%"

    _highlight_statuses(
        worksheet=worksheet,
        row_count=row_count,
    )


def _column_width(
    header: str,
) -> int:
    wide_headers = {
        "Подразделение",
        "Руководитель",
        "ФИО",
        "Должность",
        "Комментарий МПО",
        "Наставник ознакомительной стажировки",
        "Наставник основной стажировки",
        "Причина неполной оплаты",
        "Комментарий МПО 1",
        "Комментарий МПО 2",
        "Комментарий МПО 3",
        "Причина увольнения",
        "Причина увольнения — ОС руководителя",
    }

    if header in wide_headers:
        return 32

    if "Причина" in header:
        return 26

    if "Дата" in header or "срок" in header.lower():
        return 16

    if "Сумма" in header:
        return 18

    return 20


def _highlight_statuses(
    worksheet,
    row_count: int,
) -> None:
    if row_count == 0:
        return

    header_indexes = {
        cell.value: cell.column
        for cell in worksheet[1]
    }

    red_fill = PatternFill(
        fill_type="solid",
        fgColor="F4CCCC",
    )

    yellow_fill = PatternFill(
        fill_type="solid",
        fgColor="FFF2CC",
    )

    green_fill = PatternFill(
        fill_type="solid",
        fgColor="D9EAD3",
    )

    risk_headers = [
        "Риск зоны 1",
        "Риск зоны 2",
        "Риск зоны 3",
    ]

    deadline_headers = [
        "Срок 1 адаптации соблюдён",
        "Срок 2 адаптации соблюдён",
        "Срок 3 адаптации соблюдён",
    ]

    for row_index in range(
        2,
        row_count + 2,
    ):
        for header in risk_headers:
            cell = worksheet.cell(
                row=row_index,
                column=header_indexes[header],
            )

            if cell.value == "Красный":
                cell.fill = red_fill
            elif cell.value == "Жёлтый":
                cell.fill = yellow_fill
            elif cell.value == "Зелёный":
                cell.fill = green_fill

        for header in deadline_headers:
            cell = worksheet.cell(
                row=row_index,
                column=header_indexes[header],
            )

            if cell.value == "НЕТ":
                cell.fill = red_fill
            elif cell.value == "ДА":
                cell.fill = green_fill